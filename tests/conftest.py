"""Shared fixtures and explicit opt-ins for tests with external resources."""
import os
from pathlib import Path

import pytest

# Tests use explicit environment variables only, never a developer's .env file.
os.environ["PYTHON_DOTENV_DISABLED"] = "1"


OPTIONAL_TESTS = {
    "external_api": "--run-external-api",
    "local_model": "--run-local-model",
    "pdf_model": "--run-pdf-model",
}


def pytest_addoption(parser):
    for marker, option in OPTIONAL_TESTS.items():
        parser.addoption(option, action="store_true", default=False,
                         help=f"Enable {marker} tests (may use network or model downloads)")


def pytest_collection_modifyitems(config, items):
    for item in items:
        for marker, option in OPTIONAL_TESTS.items():
            if item.get_closest_marker(marker) and not config.getoption(option):
                item.add_marker(pytest.mark.skip(reason=f"Opt in with {option}"))


@pytest.fixture
def fixture_dir():
    return Path(__file__).parent / "fixtures"


@pytest.fixture
def sample_docx(tmp_path):
    from docx import Document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("RAG sample notes", level=1)
    document.add_paragraph(
        "This document is a small original fixture for the parser tests. "
        "It describes a local document collection with no private information."
    )
    document.add_paragraph(
        "A retrieval step selects relevant passages before an answer is generated. "
        "The application should preserve the source filename for each passage."
    )
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Format", "Status"), ("Markdown", "Ready"), ("HTML", "Ready")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(path)
    return path
