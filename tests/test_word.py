"""Word fixtures are generated in pytest's temporary directory."""
from io import BytesIO

from docx import Document

from src.parsers.word_parser import WordParser
from src.parsers.table_parser import TableProcessor


def test_parse_file(sample_docx):
    result = WordParser().parse(str(sample_docx))
    assert result.filename == "sample.docx"
    assert result.total_pages == 1
    assert result.total_chunks == len(result.chunks) > 0
    assert "RAG sample notes" in result.raw_text
    assert "\n\n" in result.raw_text


def test_parse_from_bytes(sample_docx):
    parser = WordParser()
    result = parser.parse_from_bytes(sample_docx.read_bytes(), sample_docx.name)
    expected = parser.parse(str(sample_docx))
    assert result.raw_text == expected.raw_text
    assert result.tables == expected.tables
    assert result.filename == sample_docx.name


def test_supported_extensions():
    assert ".docx" in WordParser().supported_extensions()


def test_table_extraction_and_chunks(sample_docx):
    result = WordParser().parse(str(sample_docx))
    assert result.table_count == 1
    assert result.tables[0]["rows"] == 3
    assert result.tables[0]["cols"] == 2
    assert "| Markdown | Ready |" in result.tables[0]["markdown"]
    table_chunks = [chunk for chunk in result.chunks if chunk["type"] == "table"]
    assert table_chunks
    assert any("Markdown" in chunk["content"] for chunk in table_chunks)


def test_table_to_dataframe(sample_docx):
    table = WordParser().parse(str(sample_docx)).tables[0]
    dataframe = TableProcessor().markdown_to_dataframe(table["markdown"])
    assert list(dataframe.columns) == ["Format", "Status"]
    assert dataframe.to_dict("records") == [
        {"Format": "Markdown", "Status": "Ready"},
        {"Format": "HTML", "Status": "Ready"},
    ]


def test_empty_document():
    stream = BytesIO()
    Document().save(stream)
    result = WordParser().parse_from_bytes(stream.getvalue(), "empty.docx")
    assert result.raw_text == ""
    assert result.chunks == []
    assert result.tables == []
